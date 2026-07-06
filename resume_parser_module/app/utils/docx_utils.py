"""
Utility helpers for extracting raw text from DOCX resumes.
"""

import io
import logging

import docx

logger = logging.getLogger(__name__)


class DOCXExtractionError(Exception):
    """Raised when text cannot be extracted from a DOCX file."""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract plain text from a DOCX file's raw bytes.

    Reads paragraph text as well as text inside tables, since resumes
    commonly use tables for layout (e.g. skills/education sections).

    Args:
        file_bytes: Raw bytes of the uploaded DOCX file.

    Returns:
        The concatenated text content of the document.

    Raises:
        DOCXExtractionError: If the file cannot be opened/parsed, or if
            no extractable text is found.
    """
    try:
        document = docx.Document(io.BytesIO(file_bytes))

        text_chunks: list[str] = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_chunks.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_chunks.append(cell.text)

        full_text = "\n".join(text_chunks).strip()

        if not full_text:
            raise DOCXExtractionError(
                "No extractable text found in DOCX file. The file may "
                "be empty or contain only images."
            )

        return full_text

    except DOCXExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - convert any python-docx error
        logger.exception("Failed to extract text from DOCX")
        raise DOCXExtractionError(f"Could not read DOCX file: {exc}") from exc
